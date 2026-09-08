"""Exportar/importar citas como planilla de Excel (.xlsx), para que el gerente
pueda revisar o editar muchas citas a la vez fuera del panel admin.

Alcance a propósito acotado: la planilla "Sesiones" es SOLO informativa (para
ver de un vistazo todo el calendario) — no se puede crear ni editar horarios
de sesión importando un Excel. Eso sigue yendo por el calendario real
(reservar.html) o el inline de SesionCita en /admin, que son los únicos
lugares que corren la validación de choques de horario (ver services.py). Si
se pudiera reprogramar sesiones por Excel, se podría saltar esa validación
sin querer con solo escribir cualquier fecha en una celda.

La planilla "Citas" sí es editable: permite crear citas nuevas o actualizar
las existentes (cliente, tatuador, estilo, sesiones, costo, estado, notas)
en lote, identificándolas por el "id" de la primera columna.
"""

import io

from django.contrib.auth import get_user_model
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.artists.models import Empleado
from apps.gallery.models import Estilo

from .models import Cita, EstadoCita

Usuario = get_user_model()

COLUMNAS_CITAS = [
    "id", "cliente_username", "tatuador_username", "estilo", "numero_sesiones",
    "costo", "estado", "notas", "fecha_finalizada",
]
COLUMNAS_SESIONES = ["cita_id", "numero", "inicio", "duracion_minutos", "estado"]


def exportar_citas_excel(queryset):
    """Arma el archivo .xlsx a partir de un queryset de Cita y lo devuelve
    como HttpResponse listo para descargar."""
    libro = Workbook()

    hoja_citas = libro.active
    hoja_citas.title = "Citas"
    hoja_citas.append(COLUMNAS_CITAS)
    for cita in queryset.select_related("cliente", "tatuador__usuario", "estilo"):
        hoja_citas.append([
            cita.pk,
            cita.cliente.username,
            cita.tatuador.usuario.username if cita.tatuador else "",
            cita.estilo.nombre if cita.estilo else "",
            cita.numero_sesiones,
            float(cita.costo) if cita.costo is not None else None,
            cita.estado,
            cita.notas,
            timezone.localtime(cita.fecha_finalizada).replace(tzinfo=None) if cita.fecha_finalizada else None,
        ])

    hoja_sesiones = libro.create_sheet("Sesiones")
    hoja_sesiones.append(COLUMNAS_SESIONES)
    for cita in queryset.prefetch_related("sesiones"):
        for sesion in cita.sesiones.all():
            hoja_sesiones.append([
                cita.pk, sesion.numero,
                timezone.localtime(sesion.inicio).replace(tzinfo=None),
                sesion.duracion_minutos, sesion.estado,
            ])

    buffer = io.BytesIO()
    libro.save(buffer)
    respuesta = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    nombre = f"citas_{timezone.localdate().isoformat()}.xlsx"
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return respuesta


def _valor(fila, indice):
    return fila[indice].value if indice < len(fila) else None


def importar_citas_excel(archivo):
    """Lee la hoja "Citas" de un .xlsx subido y crea/actualiza registros.
    Identifica la fila por la columna "id": vacía = crear una cita nueva,
    con un número = actualizar esa cita si existe. Devuelve un resumen
    (creadas, actualizadas, errores) — nunca lanza excepción por una fila
    mala, la salta y sigue con las demás."""
    try:
        libro = load_workbook(archivo, data_only=True)
    except Exception as exc:
        return {"creadas": 0, "actualizadas": 0, "errores": [f"No se pudo leer el archivo: {exc}"]}

    if "Citas" not in libro.sheetnames:
        return {"creadas": 0, "actualizadas": 0, "errores": ['El archivo no tiene una hoja llamada "Citas".']}

    hoja = libro["Citas"]
    filas = list(hoja.iter_rows())
    creadas = actualizadas = 0
    errores = []

    for n, fila in enumerate(filas[1:], start=2):  # fila 1 = encabezado
        if all(c.value in (None, "") for c in fila):
            continue  # fila vacía, se ignora

        try:
            id_crudo = _valor(fila, 0)
            cliente_username = (_valor(fila, 1) or "").strip()
            tatuador_username = (_valor(fila, 2) or "").strip()
            estilo_nombre = (_valor(fila, 3) or "").strip()
            numero_sesiones = int(_valor(fila, 4) or 1)
            costo = _valor(fila, 5)
            estado = (_valor(fila, 6) or EstadoCita.PENDIENTE).strip().upper()
            notas = _valor(fila, 7) or ""

            if not cliente_username:
                errores.append(f"Fila {n}: falta cliente_username.")
                continue
            try:
                cliente = Usuario.objects.get(username=cliente_username)
            except Usuario.DoesNotExist:
                errores.append(f"Fila {n}: no existe ningún usuario '{cliente_username}'.")
                continue

            tatuador = None
            if tatuador_username:
                tatuador = Empleado.objects.filter(usuario__username=tatuador_username).first()
                if tatuador is None:
                    errores.append(f"Fila {n}: no existe ningún tatuador '{tatuador_username}'.")
                    continue

            estilo = None
            if estilo_nombre:
                estilo = Estilo.objects.filter(nombre=estilo_nombre).first()
                if estilo is None:
                    errores.append(f"Fila {n}: no existe el estilo '{estilo_nombre}'.")
                    continue

            if estado not in EstadoCita.values:
                errores.append(f"Fila {n}: estado '{estado}' inválido (usar {', '.join(EstadoCita.values)}).")
                continue

            cita = None
            if id_crudo not in (None, ""):
                cita = Cita.objects.filter(pk=int(id_crudo)).first()
                if cita is None:
                    errores.append(f"Fila {n}: no existe ninguna cita con id {id_crudo}.")
                    continue

            if cita is None:
                cita = Cita(cliente=cliente)
                creadas += 1
            else:
                actualizadas += 1

            cita.cliente = cliente
            cita.tatuador = tatuador
            cita.estilo = estilo
            cita.numero_sesiones = numero_sesiones
            cita.costo = costo
            cita.estado = estado
            cita.notas = notas
            if estado == EstadoCita.TERMINADA and not cita.fecha_finalizada:
                cita.fecha_finalizada = timezone.now()
            cita.full_clean(exclude=["fecha_finalizada"])
            cita.save()
        except Exception as exc:
            errores.append(f"Fila {n}: {exc}")

    return {"creadas": creadas, "actualizadas": actualizadas, "errores": errores}
